import json
import re
import os
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import sys

JSON_PATH = sys.argv[1] if len(sys.argv) > 1 else "output/products_clean.json"
XLSX_PATH = sys.argv[2] if len(sys.argv) > 2 else "output/shopify_ai_sales_knowledge_base.xlsx"
CSV_PATH = sys.argv[3] if len(sys.argv) > 3 else "output/shopify_ai_sales_knowledge_base.csv"

def extract_target_age(tags, title, description):
    """Extract age specifications from tags or text."""
    age_matches = []
    
    for tag in tags:
        if re.search(r'\b(year|years|month|months|سنة|سنوات|أشهر|شهر)\b', tag, re.I):
            age_matches.append(tag.strip())
        elif re.search(r'\b\d+[-+]\d*\s*(y|yr|years?)\b', tag, re.I):
            age_matches.append(tag.strip())
            
    if age_matches:
        return "; ".join(dict.fromkeys(age_matches))
        
    title_match = re.search(r'\(Ages?\s*([^)]+)\)', title, re.I)
    if title_match:
        return title_match.group(1).strip()
        
    desc_match = re.search(r'Recommended Age:\s*([^\n•\r]+)', description, re.I)
    if desc_match:
        return desc_match.group(1).strip()
        
    return ""

def build_discount_info(variants):
    discounts = []
    for v in variants:
        price = v.get("price") or 0.0
        compare = v.get("compare_at_price")
        if compare and compare > price:
            diff = compare - price
            pct = round((diff / compare) * 100)
            discounts.append(f"Was {compare:.2f} SAR (Save {diff:.2f} SAR - {pct}%)")
            
    if discounts:
        return "; ".join(dict.fromkeys(discounts))
    return "Regular Price"

def build_price_range(variants):
    prices = [v.get("price") for v in variants if v.get("price") is not None]
    if not prices:
        return ""
    min_p, max_p = min(prices), max(prices)
    if min_p == max_p:
        return f"{min_p:.2f} SAR"
    return f"{min_p:.2f} - {max_p:.2f} SAR"

def build_stock_status(variants, product_status):
    if product_status != "ACTIVE":
        return f"INACTIVE ({product_status})"
        
    total_qty = 0
    has_trackable_qty = False
    any_available = False
    
    for v in variants:
        if v.get("available_for_sale"):
            any_available = True
        qty = v.get("inventory_quantity")
        if qty is not None:
            has_trackable_qty = True
            total_qty += max(0, qty)
            
    if not any_available or (has_trackable_qty and total_qty <= 0):
        return "OUT_OF_STOCK (0 units)"
        
    if has_trackable_qty:
        return f"IN_STOCK ({total_qty} units)"
    return "IN_STOCK (Available)"

def build_variants_summary(variants):
    summaries = []
    for v in variants:
        title = v.get("variant_title") or "Default"
        sku = v.get("sku") or "NO-SKU"
        price = v.get("price") or 0.0
        qty = v.get("inventory_quantity")
        avail = v.get("available_for_sale")
        
        stock_label = f"In Stock: {qty}" if qty is not None else ("In Stock" if avail else "Sold Out")
        if not avail and (qty is None or qty <= 0):
            stock_label = "Sold Out"
            
        summaries.append(f"[{sku}] {title}: {price:.2f} SAR ({stock_label})")
        
    return "\n".join(summaries)

def main():
    if not os.path.exists(JSON_PATH):
        raise FileNotFoundError(f"Missing {JSON_PATH}. Run 'npm run export:products' first.")

    with open(JSON_PATH, "r", encoding="utf-8") as f:
        products = json.load(f)

    print(f"Loaded {len(products)} products from {JSON_PATH}...")

    wb = Workbook()
    ws = wb.active
    ws.title = "Shopify AI Knowledge Base"
    ws.views.sheetView[0].showGridLines = True

    headers = [
        ("default_id", "Stable ID", 22),
        ("product_name", "Product Name (Title)", 36),
        ("target_age", "Target Age", 20),
        ("category_and_tags", "Category & Tags", 30),
        ("stock_status", "Stock Status", 18),
        ("price_range", "Price Range", 16),
        ("discount_info", "Discount / Promotion", 22),
        ("description", "Clean Description & Sales Copy", 55),
        ("variants_summary", "Variants (SKU, Price, Stock)", 40),
        ("image_url", "Featured Image URL", 28),
        ("product_url", "Product URL", 28),
    ]

    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    border_thin = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    fill_empty_input = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
    border_empty_input = Border(
        left=Side(style="thin", color="F59E0B"),
        right=Side(style="thin", color="F59E0B"),
        top=Side(style="thin", color="F59E0B"),
        bottom=Side(style="thin", color="F59E0B"),
    )

    for col_idx, (col_id, col_label, width) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = border_thin
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 34

    zebra_even = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    zebra_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    font_regular = Font(name="Calibri", size=10, color="1E293B")
    font_bold = Font(name="Calibri", size=10, bold=True, color="0F172A")
    font_in_stock = Font(name="Calibri", size=10, bold=True, color="166534")
    font_out_stock = Font(name="Calibri", size=10, bold=True, color="991B1B")
    fill_in_stock = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fill_out_stock = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    csv_rows = []
    csv_header = [h[0] for h in headers]
    csv_rows.append(csv_header)

    empty_cell_count = 0

    for row_idx, p in enumerate(products, 2):
        pid = p.get("numeric_product_id") or p.get("shopify_product_id", "").split("/")[-1]
        default_id = f"prod-{pid}"
        title = p.get("title") or ""
        desc = p.get("description") or ""
        category = p.get("category") or p.get("product_type") or "Uncategorized"
        tags = p.get("tags") or []
        variants = p.get("variants") or []
        status = p.get("status") or "ACTIVE"

        target_age = extract_target_age(tags, title, desc)
        tags_str = ", ".join(tags[:6])
        category_and_tags = f"Category: {category} | Tags: {tags_str}" if tags_str else f"Category: {category}"
        stock_status = build_stock_status(variants, status)
        price_range = build_price_range(variants)
        discount_info = build_discount_info(variants)
        variants_summary = build_variants_summary(variants)
        image_url = p.get("featured_image_url") or ""
        product_url = p.get("product_url") or ""

        row_data = [
            default_id,
            title,
            target_age,
            category_and_tags,
            stock_status,
            price_range,
            discount_info,
            desc,
            variants_summary,
            image_url,
            product_url,
        ]

        csv_rows.append(row_data)

        row_fill = zebra_even if row_idx % 2 == 0 else zebra_odd

        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = border_thin
            cell.fill = row_fill
            cell.font = font_regular

            is_empty = val is None or str(val).strip() == ""

            if is_empty:
                cell.fill = fill_empty_input
                cell.border = border_empty_input
                empty_cell_count += 1
            else:
                if col_idx in [1, 3, 6]:
                    cell.alignment = Alignment(horizontal="center", vertical="top")
                    if col_idx == 1:
                        cell.font = font_bold
                elif col_idx in [8, 9, 4]:
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                elif col_idx == 5:
                    cell.alignment = Alignment(horizontal="center", vertical="top")
                    if "IN_STOCK" in str(val):
                        cell.fill = fill_in_stock
                        cell.font = font_in_stock
                    else:
                        cell.fill = fill_out_stock
                        cell.font = font_out_stock
                elif col_idx in [10, 11]:
                    cell.alignment = Alignment(horizontal="left", vertical="top")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="top")

        ws.row_dimensions[row_idx].height = 45

    ws.freeze_panes = "A2"

    wb.save(XLSX_PATH)
    print(f"Successfully saved Excel: {XLSX_PATH} ({os.path.getsize(XLSX_PATH)/1024:.1f} KB)")
    print(f"Highlighted {empty_cell_count} empty cells in yellow input styling.")

    with open(CSV_PATH, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        for r in csv_rows:
            writer.writerow(r)
    print(f"Successfully saved CSV: {CSV_PATH} ({os.path.getsize(CSV_PATH)/1024:.1f} KB)")

if __name__ == "__main__":
    main()
